"""
Idle Clans - Gold/h Farming-Analyse

Zieht Marktpreise + Rezept-Daten aus der Idle-Clans-API, rechnet fuer jedes farmbare
Item Gold/h unter Beruecksichtigung deiner Account-Upgrades (Speed/Yield/Cost-Boosts,
Handschuhe, XP-Boosts, NPC-Verkaufspreis) und exportiert das Ergebnis nach Excel.

EIGENSTAENDIGES SKRIPT: laeuft unabhaengig vom Autoclicker (kein Import aus `autoclicker/`,
kein Windows noetig). Zusaetzliche Abhaengigkeiten, die der Autoclicker selbst NICHT
braucht - deshalb bewusst nicht in requirements.txt:

    pip install pandas requests openpyxl matplotlib

(matplotlib nur fuer SHOW_PRICE_SENSITIVITY_CHART - fehlt es, laeuft der Rest weiter.)

Aufruf:  python tools/market_analysis.py

Struktur:
  1. API-Endpunkte
  2. Account-Konfiguration (SKILLS) - alles Spieler-Spezifische an einem Ort
  3. Markt-Filter-Schwellenwerte
  4. Daten laden (API)
  5. Markt-/Preis-Hilfsfunktionen (inkl. NPC-Verkaufspreis-Fallback)
  6. Rezept-Normalisierung (Upgrades + Best-/Worst-Case-Kosten)
  7. Einzelschritt-Analyse ("Rohdaten"-Tab)
  8. Ketten-Analyse ("Ketten"/"Realistisch_Farmbar"-Tabs) + Preis-Sensitivitaets-Chart
  9. Excel-Export
  10. Lauf-Sanity-Check (Vergleich mit letztem Lauf)
  11. main()

v12-Aenderungen (Review-Durchgang, Details im jeweiligen Kommentar am Code):
  - API-Requests laufen ueber _get_json(): HTTP-Status wird geprueft, statt bei einem
    500er/HTML-Fehlerdokument mit einem kryptischen JSONDecodeError abzustuerzen.
    Zusaetzlich Schema-Drift-Warnungen (leere Antwort / dailyAveragePrice nirgends
    gesetzt / unplausible Aktionszeiten).
  - Plausibilitaetscheck fuer die Zeiteinheit von BaseTime (Annahme: Millisekunden).
    Kippt die API mal auf Sekunden, ist sonst JEDE Gold/h-Zahl um Faktor 1000 daneben,
    ohne dass irgendwas auffaellt.
  - Smelting Magic gilt lt. Wiki NICHT fuer Astronomical ore -> astronomical_bar bekommt
    den Rabatt nicht mehr (SMELTING_MAGIC_EXCLUDED_SUBSTRINGS).
  - Preis-Sensitivitaet rechnet jetzt NETTO (Gold/s minus Materialkosten), wie alle
    anderen Gold/h-Spalten auch. Vorher war es Brutto-Umsatz und damit nicht mit dem
    Gold/h aus Ketten/Realistisch_Farmbar vergleichbar.
  - Worst-Case-Merges: how="left" + validate="one_to_one" statt Inner-Join. Ein
    doppelter (ItemID, TaskId)-Schluessel haette sonst still Zeilen vervielfacht bzw.
    verschluckt, statt einen Fehler zu werfen.
  - matplotlib wird erst beim Chart importiert UND der ImportError abgefangen - vorher
    ist der komplette Lauf (inkl. Excel-Export) daran gestorben, dass ein reines
    Komfort-Feature nicht installiert war.
  - run_stats_history.json fuehrt jetzt tatsaechlich eine Historie (Liste der letzten
    RUN_STATS_HISTORY_LIMIT Laeufe) statt nur den letzten Lauf; altes Dict-Format wird
    weiterhin gelesen.
  - "Gold pro Stück" ohne max(..., 1)-Klemme (verfaelschte bei <1 Stueck/h das Ergebnis).
  - Guard gegen speed_factor <= 0 (Equipment-Boost >= 100% haette eine Division durch
    0 bzw. negative Zeiten erzeugt).

v10-Aenderungen:
  - Chart-Y-Achse auf Gold/s umgestellt (uebersichtlicher als Gold/h bei vielen Linien).
  - Rohdaten hat keinen Hard-Filter mehr: JEDES Recipe landet in der Tabelle, auch wenn
    es nicht in Ketten/Realistisch_Farmbar/Nach_Skill_Level auftaucht (die bleiben wie
    gehabt gefiltert - "nur was Sinn macht"). Neue Spalten "InKetten"/"Status"/
    "AusschlussGrund" zeigen, warum. Excel-Faerbung: ganze Zeile kraeftig orange
    (F6B26B), das verantwortliche Feld "AusschlussGrund" zusaetzlich kraeftig rot
    (E06666). Gold/h behaelt sein gelbes Sortier-Highlight.
  - Neues Sheet "Preis_Sensitivitaet": dieselben Live-Orderbook-Daten, die auch in den
    PNG-Chart einfliessen (SHOW_PRICE_SENSITIVITY_CHART), zusaetzlich tabellarisch.
  - Neuer optionaler SHOW_LONGTERM_AVERAGES-Flag: ergaenzt Rohdaten um Avg1D/Avg7D/
    Avg30D/Volume1D aus dem comprehensive-Endpoint (1 Request pro eindeutigem Item -
    kann dauern). Feldnamen sind ein Best-Guess, s. COMPREHENSIVE_AVG_FIELDS.
  - Neuer Lauf-Sanity-Check (Abschnitt 10): vergleicht Tasks-/Rezept-/Item-Anzahl sowie
    Markt-Kennzahlen mit dem letzten Lauf (run_stats_history.json) und warnt bei
    verdaechtigen Ausschlaegen, die eher auf ein kaputtes API-Feld als auf echte
    Marktbewegung hindeuten.

v9-Aenderungen:
  - Best-/Worst-Case fuer "Smelting Magic" (30% erwartete Ore->Bar-Ersparnis): der Wiki-
    Kommentar spricht nur von "Erz", das Script wendet den Rabatt bisher (Best-Case) auf
    ALLE Cost-Zeilen der *_bar-Rezepte an (also auch auf z.B. Kohle). Worst-Case: Rabatt
    gilt nur auf die erste Cost-Zeile (Haupt-Erz), Nebenzutaten voller Preis. Beide
    Varianten werden jetzt PARALLEL durchgerechnet und als "_Worst"-Spalten in Rohdaten
    UND Ketten/Realistisch_Farmbar ausgegeben (CostCaseAmbiguous=True markiert betroffene
    Zeilen - betrifft nur *_bar-Rezepte selbst und alles, was direkt/indirekt davon
    abhaengt, z.B. titanium_platebody haengt an titanium_bar).
  - Neuer optionaler Preis-Sensitivitaets-Chart (SHOW_PRICE_SENSITIVITY_CHART, Standard
    aus): matplotlib-Liniendiagramm ueber die Top-N FullySelfSufficient-Items (alle
    Skills, wie Realistisch_Farmbar) mit Gold/s auf einer gemeinsamen Y-Achse ueber die
    10 aktuellen Orderbook-Preispunkte (5 Buy aufsteigend, 5 Sell aufsteigend) je Item.
    Nutzt die bereits vorhandene fetch_orderbook_depth()-Funktion (Live-Requests, s. Flag).

v8-Aenderungen (per Ingame-Boosts-Breakdown verifiziert):
  - Farming ist laut Wiki ein Processing-Skill, KEIN Gathering-Skill -> hat keinen
    "Gatherers"-Perk (5% Clan-Speed). is_gathering wurde faelschlich auf True gesetzt.
  - XP_BOOST_TOTAL war 25%+30% (Wiki-Beispielwerte), tatsaechlich aktiv sind laut
    Boosts-Breakdown nur Clan house (25%) + House (25%) = 50%, kein Castle-Tier.
  - NPC-Verkaufspreis (base_value aus Items-API) wird jetzt als Verkaufsweg
    beruecksichtigt: manche Items haben keinen/kaum liquiden Player-Markt, sind aber
    per Sofortverkauf an den NPC-Shop trotzdem brauchbar. "An offer they can't refuse"
    (Clan-Upgrade, einmalig, kein Tier-System) gibt +10% auf diesen NPC-Preis.
"""

from __future__ import annotations

import json
import re
import time
from dataclasses import dataclass
from datetime import datetime

import pandas as pd
import requests
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# 1. API-ENDPUNKTE
# ============================================================

# Hinweis: es gibt zusaetzlich /PlayerMarket/items/prices/latest/all. Falls der
# Endpunkt unten irgendwann nur noch ein Teilset liefert (s. Sanity-Check-Warnung
# "market_map_count eingebrochen"), ist /latest/all die erste Alternative zum Testen.
MARKET_URL = "https://query.idleclans.com/api/PlayerMarket/items/prices/latest?includeAveragePrice=true"
GAME_URL = "https://query.idleclans.com/api/Configuration/game-data"
COMPREHENSIVE_URL_TEMPLATE = "https://query.idleclans.com/api/PlayerMarket/items/prices/latest/comprehensive/{item_id}"

# ============================================================
# 2. ACCOUNT-KONFIGURATION
#
# Alles, was von deinem Account/Equipment abhaengt, steht hier an EINER Stelle statt
# verstreut. equipment_speed_boost = "Equipment skill boosts" aus dem Ingame Boosts-
# Breakdown-Screen. is_gathering = bekommt den Gatherers-Clan-Bonus (+5% Speed) - NUR
# echte Gathering-Skills (Mining/Fishing/Foraging/Woodcutting), NICHT Farming (laut
# Wiki ein Processing-Skill, kein Gatherers-Perk vorhanden). gloves_owned = du besitzt
# die Skilling-Handschuhe fuer diesen Skill (5% Chance auf doppelte Beute, stackt
# multiplikativ mit yield_multiplier). yield_multiplier = fixe Verdopplungs-Upgrades
# (Fisherman/Lumberjack/Power Forager). cost_multiplier = Anteil der Materialkosten,
# der nach Trickery/Magic-Ersparnis noch anfaellt (uniform auf alle Cost-Zeilen - fuer
# Smithing-*_bar-Rezepte gilt stattdessen die Best-/Worst-Case-Logik in normalize_recipe,
# s. SMITHING_SMELTING_COST_MULTIPLIER). excluded = Skill komplett ignorieren
# (Combat/Enchanting/etc., kein normales Markt-Item-Recipe).
#
# OFFENE ANNAHME (bewusst nicht "korrigiert", weil nur ingame messbar): das Script
# verrechnet Clan-Speed und Equipment-Speed MULTIPLIKATIV auf die Aktionszeit
# (t * (1-0.05) * (1-0.61)). Rechnet das Spiel stattdessen additiv (t * (1-0.66)),
# faellt jede Gold/h-Zahl ca. 9% zu niedrig aus - die RANGFOLGE der Items bleibt aber
# gleich, weil der Faktor auf alle Rezepte desselben Skills identisch wirkt.
# Gegenprobe: eine bekannte Aktion ingame stoppen und mit Time_sec in Rohdaten
# vergleichen; weicht es ab, hier auf die additive Variante umstellen.
# ============================================================


@dataclass
class SkillConfig:
    equipment_speed_boost: float = 0.0
    is_gathering: bool = False
    gloves_owned: bool = False
    yield_multiplier: float = 1.0
    cost_multiplier: float = 1.0
    excluded: bool = False


GLOVES_DOUBLE_CHANCE = 0.05          # 5% Chance auf doppelte Beute, alle Skilling-Handschuhe
CLAN_GATHERERS_SPEED_BOOST = 0.05    # Clan-Upgrade "Gatherers", nur fuer is_gathering=True

# XP-Boost: verifiziert per Ingame Boosts-Breakdown-Screen. Aktiv sind "Clan house"
# (Clan-Upgrade, 25%) und "House" (persoenliches Upgrade, 25%). Kein Castle-Tier
# vorhanden - falls spaeter eins dazukommt, hier ergaenzen.
XP_BOOST_TOTAL = 0.25 + 0.25

# Daily Boost ("Flamme" im Profil, alle 2h manuell aktivierbar, mit Premium-Token alle
# 8h) - PER WIKI-FORMEL MULTIPLIKATIV mit XP_BOOST_TOTAL verknuepft, nicht additiv:
#   XP = Base * (1 + Clan house + House) * (1 + Daily Boost) * (1 + Jewelry/Equip-Boost)
# Per Live-Abgleich (Coal ore, Sunflowerberry) verifiziert: der Boost gibt +4%. Ist
# TEMPORAER (laeuft nach einigen Stunden ab) - deshalb als An/Aus-Flag statt Rohwert:
# vor dem Ausfuehren kurz pruefen, ob die Flamme im Profil gerade aktiv ist.
#
# ACHTUNG (Review): das Wiki nennt fuer den Daily Boost +30% XP fuer 2h (bis 4x/Tag,
# max. 8h). Die hier gemessenen +4% passen nicht dazu - entweder wurde etwas anderes
# gemessen (z.B. ein Jewelry-Boost) oder der Wert wurde im Spiel geaendert. Der Boost
# beeinflusst NUR die XP-Spalten (XP/h, Gold per XP), kein Gold/h - deshalb hier nur
# als Notiz und nicht "einfach korrigiert". Bei Gelegenheit gegenmessen.
DAILY_BOOST_ACTIVE = False   # True setzen, wenn die Flamme im Profil gerade aktiv ist
DAILY_XP_BOOST_PERCENT = 0.04
DAILY_XP_BOOST = DAILY_XP_BOOST_PERCENT if DAILY_BOOST_ACTIVE else 0.0

# NPC-Verkaufspreis-Boost: zwei permanente Quellen, multiplikativ kombiniert.
# Per Wiki bestaetigt: +10% (Clan) kombiniert mit +5% (Potion of negotiation) ergibt
# exakt 1.155x Vendor-Value - also multiplikativ, genau wie hier gerechnet.
OFFER_THEY_CANT_REFUSE_ACTIVE = True   # Clan-Upgrade, einmalig, +10%
POTION_OF_NEGOTIATION_ACTIVE = True    # bei Timo als PERMANENT bestaetigt, +5%
NPC_SELL_BOOST_MULTIPLIER = (
    (1.10 if OFFER_THEY_CANT_REFUSE_ACTIVE else 1.0)
    * (1.05 if POTION_OF_NEGOTIATION_ACTIVE else 1.0)
)

AUTO_COOK_CHANCE = 0.5               # 50% der gefangenen Fische sind beim Fishing bereits gekocht
AUTO_COOK_SOURCE_SKILL = "Fishing"

# Smelting Magic (Wiki bestaetigt): Chance, Erz beim SCHMELZEN (Ore -> Bar) nicht zu
# verbrauchen. Gilt NICHT fuers Schmieden von Bars zu Ruestung (war vorher ein Bug: der
# Rabatt wurde faelschlich auf ALLE Smithing-Rezepte angewendet). Wird unten in
# normalize_recipe() nur bei Rezepten angewendet, deren Name auf "_bar" endet.
#
# UNKLARE REICHWEITE (v9): der Wiki-Text nennt nur "Erz" - bei *_bar-Rezepten mit
# mehreren Cost-Zeilen (z.B. titanium_bar = Erz + Kohle) ist unklar, ob der Rabatt auf
# ALLE Zeilen wirkt oder nur auf die erste (das Haupt-Erz). Best-Case = alle Zeilen
# (bisherige v8-Annahme), Worst-Case = nur die erste Zeile. Beide werden parallel
# gerechnet, s. normalize_recipe(case=...) und die "_Worst"-Spalten in Rohdaten/Ketten.
#
# Tier-System: die Tiers stacken NICHT, der neueste Tier ueberschreibt den vorherigen -
# 0.3 entspricht dem hoechsten Tier. Bei niedrigerem Tier hier anpassen.
SMITHING_SMELTING_COST_MULTIPLIER = 1.0 - 0.3

# Ausnahme lt. Wiki/Community-Guide: Astronomical ore ist vom Smelting-Magic-Effekt
# ausgenommen. Der Abgleich laeuft ueber den Rezeptnamen (klein, Teilstring), weil das
# Script nur den Namen kennt - falls die API das Rezept anders benennt, hier ergaenzen.
SMELTING_MAGIC_EXCLUDED_SUBSTRINGS = ("astronomical",)

SKILLS: dict[str, SkillConfig] = {
    "Mining":      SkillConfig(equipment_speed_boost=0.55 + 0.06, is_gathering=True, gloves_owned=True),
    "Fishing":     SkillConfig(equipment_speed_boost=0.55 + 0.06, is_gathering=True, gloves_owned=True,
                                yield_multiplier=2.0),   # Fisherman: 100% doppelte Ausbeute
    "Foraging":    SkillConfig(equipment_speed_boost=0.55 + 0.06, is_gathering=True, gloves_owned=True,
                                yield_multiplier=1.5),   # Power Forager: 50% Chance auf doppelte Beute
    "Woodcutting": SkillConfig(equipment_speed_boost=0.55 + 0.06, is_gathering=True, gloves_owned=True,
                                yield_multiplier=2.0),   # Lumberjack: 100% doppelte Ausbeute
    "Cooking":     SkillConfig(equipment_speed_boost=0.55 + 0.06, gloves_owned=True),
    "Carpentry":   SkillConfig(equipment_speed_boost=0.55, gloves_owned=True),
    "Smithing":    SkillConfig(equipment_speed_boost=0.55),  # Smelting Magic s. SMITHING_SMELTING_COST_MULTIPLIER oben, nicht hier
    # Farming: KEIN Gatherers-Perk (Processing-Skill lt. Wiki, is_gathering=False).
    # Farming trickery Tier 5 (hoechster Tier) = 50% Chance, Saatgut zu sparen.
    "Farming":     SkillConfig(equipment_speed_boost=0.55, is_gathering=False, cost_multiplier=0.5),
    "Crafting":    SkillConfig(equipment_speed_boost=0.55 + 0.08),
    "Agility":     SkillConfig(equipment_speed_boost=0.55 + 0.06),  # produziert Samen fuer Beeren (kein Gathering-Skill lt. Wiki)
    "Plundering":  SkillConfig(equipment_speed_boost=0.55 + 0.06, gloves_owned=True),  # Ghostly-Outfit (3-teilig) bestaetigt
    "Brewing":     SkillConfig(),  # TODO(Timo): equipment_speed_boost aus Boosts-Screen eintragen

    # Skills ohne normales Markt-Item-Recipe -> komplett ausgeschlossen
    "Attack": SkillConfig(excluded=True),
    "Strength": SkillConfig(excluded=True),
    "Defence": SkillConfig(excluded=True),
    "Archery": SkillConfig(excluded=True),
    "Magic": SkillConfig(excluded=True),
    "Health": SkillConfig(excluded=True),
    "Exterminating": SkillConfig(excluded=True),
    "Enchanting": SkillConfig(excluded=True),
    "Invocation": SkillConfig(excluded=True),
    "Item creation": SkillConfig(excluded=True),  # Key-Name gegen echte API noch unbestaetigt
}

DEFAULT_SKILL_CONFIG = SkillConfig()  # Fallback fuer unbekannte/neue Skills: konservativ, nichts ausgeschlossen


def skill_cfg(skill_name: str) -> SkillConfig:
    return SKILLS.get(skill_name, DEFAULT_SKILL_CONFIG)


# ============================================================
# 3. MARKT-FILTER (Account-unabhaengig)
# ============================================================

MIN_SELL_VOLUME = 10000        # Mindest-BuyVol des Endprodukts (Nachfrage fuer Sofortverkauf), gilt nur wenn ueber Player-Markt verkauft wird
MIN_MARKET_VOLUME = 50         # Mindest-Volumen je Seite, damit ein Markt als "echt gehandelt" gilt
MAX_SPREAD_RATIO = 1.0         # Ask darf hoechstens 2x Bid sein
MAX_AVG_DEVIATION_RATIO = 0.5  # Warnung, wenn Bid/Ask >50% vom 24h-Avg abweicht
LIQUIDITY_WARNING_RATIO = 5.0  # Bedarf/Absatz > 5x verfuegbares Volumen -> unrealistisch

TOP_N_RANKING = 25
DEPTH_ANALYSIS_TOP_N = 25
SHOW_ORDERBOOK_DEPTH = False    # True = zusaetzlich 25 Live-Requests fuer Orderbook-Tiefe (langsam)

# Preis-Sensitivitaets-Chart (v9): matplotlib-Liniendiagramm, Gold/s ueber die 10
# aktuellen Orderbook-Preispunkte (5 Buy + 5 Sell) je Item, Top-N FullySelfSufficient-
# Items (alle Skills). Zusaetzliche Live-Requests wie SHOW_ORDERBOOK_DEPTH.
SHOW_PRICE_SENSITIVITY_CHART = True
PRICE_SENSITIVITY_TOP_N = 10
PRICE_SENSITIVITY_CHART_PATH = "price_sensitivity_chart.png"

# 1-/7-/30-Tage-Durchschnittspreis + Tagesvolumen: kommen laut API-Doku aus dem
# comprehensive-Endpoint (derselbe, der auch fuer Orderbook-Tiefe/Preis-Sensitivitaet
# genutzt wird), sind aber bisher nicht typisiert/verifiziert. Feldnamen hier sind ein
# informierter Best-Guess nach dem Muster des Bulk-Endpoints ("dailyAveragePrice").
# -> Meldet die Konsole beim Lauf "Erwartete Avg-Felder nicht gefunden", stehen die
#    tatsaechlichen Keys direkt in derselben Meldung; hier eintragen.
COMPREHENSIVE_AVG_FIELDS = {
    "Avg1D": "dailyAveragePrice",
    "Avg7D": "weeklyAveragePrice",
    "Avg30D": "monthlyAveragePrice",
}
COMPREHENSIVE_VOLUME_FIELD = "dailyVolume"

# True = zusaetzlich 1 Request PRO EINDEUTIGEM ITEM in Rohdaten (koennen mehrere Hundert
# sein -> mehrere Minuten Laufzeit). Ergaenzt Avg1D/Avg7D/Avg30D/Volume1D in Rohdaten.
SHOW_LONGTERM_AVERAGES = False
LONGTERM_AVERAGES_REQUEST_DELAY_S = 0.05  # kleine Pause zwischen Requests, um die API zu schonen

EXPORT_PATH = "idle_clans_farming.xlsx"


# ============================================================
# 4. DATEN LADEN
# ============================================================

def _get_json(url: str, timeout: int = 30, context: str = "API") -> requests.Response:
    """Ein Request, ein Fehlerbild: Netzwerkfehler und HTTP-Fehlercodes werden hier zu
    einer verstaendlichen Meldung. Ohne das kippt ein 500er/HTML-Fehlerdokument spaeter
    als kryptischer JSONDecodeError raus - oder schlimmer: der Lauf rechnet mit einer
    halben Antwort weiter."""
    try:
        resp = requests.get(url, timeout=timeout)
    except requests.RequestException as exc:
        raise RuntimeError(f"{context}: Request fehlgeschlagen ({exc}) - URL: {url}") from exc
    if resp.status_code != 200:
        raise RuntimeError(f"{context}: HTTP {resp.status_code} von {url}")
    return resp


def load_market_map() -> dict:
    """ItemId -> Bid/Ask/Volumen. Bid ("buy") = Preis bei Sofortverkauf, Ask ("sell") =
    Preis bei Sofortkauf. buyVol begrenzt Sofortverkauf, sellVol begrenzt Sofortkauf."""
    resp = _get_json(MARKET_URL, timeout=30, context="Markt-Endpoint")
    payload = resp.json()
    if not isinstance(payload, list) or not payload:
        raise RuntimeError(f"Markt-Endpoint: unerwartete/leere Antwort ({type(payload).__name__}) - URL: {MARKET_URL}")

    market_map = {}
    for i in payload:
        market_map[i.get("itemId")] = {
            "buy": i.get("highestBuyPrice", 0),
            "sell": i.get("lowestSellPrice", 0),
            "buyVol": i.get("highestPriceVolume", 0),
            "sellVol": i.get("lowestPriceVolume", 0),
            "avg": i.get("dailyAveragePrice", 0),
        }

    # Schema-Drift-Warnung: die Ø-Preis-Spalten (und price_anomaly()) haengen komplett an
    # dailyAveragePrice. Wenn das Feld verschwindet/umbenannt wird, wuerde das sonst nur
    # als "keine Anomalien gefunden" durchrutschen statt als Fehler aufzufallen.
    if not any(entry["avg"] for entry in market_map.values()):
        example_keys = sorted(payload[0].keys())
        print("⚠ Markt-Endpoint: 'dailyAveragePrice' ist bei KEINEM Item gesetzt. Entweder war "
              "includeAveragePrice wirkungslos oder das Feld heisst jetzt anders.\n"
              f"  Vorhandene Keys im ersten Eintrag: {example_keys}")

    return market_map


def load_game_data() -> dict:
    """Game-Data-Endpoint liefert MongoDB-Extended-JSON (ObjectId(...) ist kein valides
    JSON) - wird vor dem Parsen bereinigt."""
    resp = _get_json(GAME_URL, timeout=60, context="Game-Data-Endpoint")
    raw = re.sub(r'ObjectId\("([a-f0-9]+)"\)', r'"\1"', resp.text)
    try:
        return json.loads(raw)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Game-Data-Endpoint: Antwort ist kein gueltiges JSON ({exc}) - "
                           f"evtl. neues Extended-JSON-Konstrukt neben ObjectId(...).") from exc


def build_item_info_map(game: dict) -> dict:
    """ItemId -> {name, base_value}. Eintraege ohne ItemId werden uebersprungen statt
    den Lauf mit einem KeyError abzubrechen (die API hat schon Platzhalter-Eintraege
    ohne alle Felder ausgeliefert)."""
    info = {}
    for it in game.get("Items", {}).get("Items", []):
        item_id = it.get("ItemId")
        if item_id is None:
            continue
        info[item_id] = {"name": it.get("Name", f"item_{item_id}"), "base_value": it.get("BaseValue", 0)}
    return info


# ============================================================
# 5. MARKT-/PREIS-HILFSFUNKTIONEN
# ============================================================

def valid_market(item: dict) -> bool:
    """Hard-Filter: hat das Item ueberhaupt einen echten, liquiden Markt? Spread ist
    bewusst KEIN Ausschlusskriterium mehr (siehe wide_spread()) - manche Items haben
    legitim einen breiten Spread (tiefer Markt), sind aber ueber eine Sell-Order statt
    Sofortverkauf trotzdem gut handelbar."""
    buy, sell = item["buy"], item["sell"]
    if buy <= 0 or sell <= 0:
        return False
    if item["buyVol"] < MIN_MARKET_VOLUME or item["sellVol"] < MIN_MARKET_VOLUME:
        return False
    return True


def wide_spread(item: dict) -> bool:
    """Weiche WARNUNG: Ask (sell) liegt um mehr als MAX_SPREAD_RATIO ueber Bid (buy).
    Kein Ausschlusskriterium - nur ein Hinweis, dass der Markt tief/duenn ist."""
    buy, sell = item["buy"], item["sell"]
    if buy <= 0:
        return False
    return abs(sell - buy) / buy > MAX_SPREAD_RATIO


def price_anomaly(item: dict) -> bool:
    """Weiche WARNUNG (kein Hard-Filter): weicht Bid/Ask stark vom 24h-Avg ab?"""
    avg = item.get("avg", 0)
    if avg <= 0:
        return False
    return abs(item["buy"] - avg) / avg > MAX_AVG_DEVIATION_RATIO or abs(item["sell"] - avg) / avg > MAX_AVG_DEVIATION_RATIO


def npc_sell_price(item_id: int, item_info_map: dict) -> float:
    """Preis bei Sofortverkauf an den NPC-Vendor (unbegrenztes Volumen, im Gegensatz
    zum Player-Markt). base_value kommt aus der Items-API, Boost durch das permanente
    Clan-Upgrade "An offer they can't refuse" (+10%) plus Potion of negotiation (+5%),
    multiplikativ = 1.155x (per Wiki bestaetigt, siehe NPC_SELL_BOOST_MULTIPLIER oben)."""
    return item_info_map.get(item_id, {}).get("base_value", 0) * NPC_SELL_BOOST_MULTIPLIER


def effective_sell_price(item_id: int, market_map: dict, item_info_map: dict) -> tuple[float, bool]:
    """Bester Verkaufsweg fuer das Endprodukt: Player-Market-Bid (nur wenn liquide genug,
    sonst Verlustgefahr durch fehlende Abnahme) vs. NPC-Vendor-Preis. Gibt
    (Preis, sold_to_npc) zurueck. sold_to_npc=True heisst: NPC ist die bessere/einzige Option."""
    m = market_map.get(item_id)
    market_price = m["buy"] if (m is not None and valid_market(m) and m["buyVol"] >= MIN_SELL_VOLUME) else 0.0
    npc_price = npc_sell_price(item_id, item_info_map)
    if npc_price > market_price:
        return npc_price, True
    return market_price, False


def is_player_shop_tradeable(item_info: dict) -> bool:
    """STUB: soll pruefen, ob ein Item im Player Market gelistet werden kann. Feldname
    dafuer im rohen Item-JSON ist unbekannt (kein Live-Zugriff moeglich) - aktuell NO-OP.
    In der Praxis filtert valid_market() (kein/kein liquider Markteintrag) das meiste
    davon bereits automatisch raus, und der NPC-Preis-Fallback faengt den Rest ab."""
    return True


def is_raid_recipe(name: str) -> bool:
    return "raids_" in str(name).lower()


# ============================================================
# 6. REZEPT-NORMALISIERUNG
#
# Jedes rohe Recipe wird GENAU EINMAL pro Case (best/worst) hier verarbeitet: Upgrades
# anwenden (Speed/Yield/Cost/XP), ungueltige Eintraege aussortieren. Alle nachgelagerten
# Analysen (Einzelschritt UND Ketten) lesen aus derselben Liste - keine doppelt
# gepflegte Logik. Fuer die Best-/Worst-Case-Unterscheidung (Smelting-Magic-Reichweite)
# wird build_all_recipes() zweimal aufgerufen (case="best"/"worst"), s. main().
# ============================================================

def _is_smelting_magic_recipe(skill_name: str, recipe_name: str) -> bool:
    """Ore -> Bar Schmelzen, auf das Smelting Magic wirkt. Ausnahmen (Astronomical ore)
    siehe SMELTING_MAGIC_EXCLUDED_SUBSTRINGS."""
    if skill_name != "Smithing" or not recipe_name.endswith("_bar"):
        return False
    lowered = recipe_name.lower()
    return not any(excl in lowered for excl in SMELTING_MAGIC_EXCLUDED_SUBSTRINGS)


def normalize_recipe(skill_name: str, raw_recipe: dict, case: str = "best") -> dict | None:
    if raw_recipe.get("Disabled", False):
        return None  # z.B. Citadel-Raid-only-Content (bestaetigt per Live-Check: Disabled=True)

    if is_raid_recipe(raw_recipe.get("Name", "")):
        return None

    base_time = raw_recipe.get("BaseTime", 0.0)
    if base_time <= 0:
        return None  # Kategorie-Platzhalter, keine echte Aktion

    item_id = raw_recipe.get("ItemReward", -1)
    item_amount = raw_recipe.get("ItemAmount", 0)
    if item_id is None or item_id < 0 or item_amount <= 0:
        return None  # kein direkter Item-Reward (z.B. Scrolls)

    cfg = skill_cfg(skill_name)
    clan_boost = CLAN_GATHERERS_SPEED_BOOST if cfg.is_gathering else 0.0
    speed_factor = (1.0 - clan_boost) * (1.0 - cfg.equipment_speed_boost)
    if speed_factor <= 0:
        # Ein Speed-Boost von >=100% ist keine gueltige Konfiguration (Aktionszeit 0 oder
        # negativ). Lieber laut und einmalig scheitern als stillschweigend Gold/h=inf.
        raise ValueError(
            f"Skill '{skill_name}': equipment_speed_boost={cfg.equipment_speed_boost} ergibt "
            f"Aktionszeit <= 0. Wert in SKILLS pruefen (0.61 = 61% schneller, nicht 61)."
        )
    yield_factor = cfg.yield_multiplier * (1.0 + GLOVES_DOUBLE_CHANCE if cfg.gloves_owned else 1.0)

    # Smelting Magic wirkt nur beim Ore->Bar-Schmelzen (Recipe-Name endet auf "_bar"),
    # nicht generell auf alle Smithing-Rezepte (siehe Kommentar bei SMITHING_SMELTING_COST_MULTIPLIER).
    # Reichweite innerhalb des Rezepts ist unklar (s. Kommentar oben) -> pro Cost-Zeile:
    # Best-Case = Rabatt auf ALLE Zeilen, Worst-Case = nur auf die erste Zeile (Haupt-Erz).
    is_bar_smelt = _is_smelting_magic_recipe(skill_name, raw_recipe.get("Name", ""))
    costs = []
    for i, c in enumerate(raw_recipe.get("Costs") or []):
        if is_bar_smelt:
            line_mult = SMITHING_SMELTING_COST_MULTIPLIER if (case == "best" or i == 0) else 1.0
        else:
            line_mult = cfg.cost_multiplier
        costs.append({"Item": c.get("Item"), "Amount": c.get("Amount", 0) * line_mult})

    return {
        "name": raw_recipe.get("Name", "unknown"),
        "skill": skill_name,
        "item_id": item_id,
        "base_time_ms": base_time * speed_factor,
        "item_amount": item_amount * yield_factor,
        "xp": raw_recipe.get("ExpReward", 0.0) * (1.0 + XP_BOOST_TOTAL) * (1.0 + DAILY_XP_BOOST),
        "level": raw_recipe.get("LevelRequirement"),
        "costs": costs,
        "task_id": raw_recipe.get("TaskId"),
        "cost_case_ambiguous": is_bar_smelt and len(costs) > 1,
    }


def build_all_recipes(tasks: dict, case: str = "best") -> list:
    all_recipes = []
    for skill_name, blocks in tasks.items():
        if skill_cfg(skill_name).excluded:
            continue
        for block in blocks:
            for raw_recipe in block.get("Items", []):
                normalized = normalize_recipe(skill_name, raw_recipe, case=case)
                if normalized is not None:
                    all_recipes.append(normalized)
    return all_recipes


# Plausibilitaetsfenster fuer eine einzelne Skilling-Aktion (nach Speed-Boosts).
# Das ganze Script rechnet mit 3_600_000 ms/h, setzt also voraus, dass BaseTime aus der
# API in MILLISEKUNDEN kommt. Waere es Sekunden, saehe man das nirgends - alle Gold/h
# waeren einfach um Faktor 1000 daneben. Deshalb ein billiger Median-Check.
MIN_PLAUSIBLE_ACTION_SEC = 0.3
MAX_PLAUSIBLE_ACTION_SEC = 300.0


def check_action_time_plausibility(all_recipes: list):
    if not all_recipes:
        return
    times = sorted(r["base_time_ms"] / 1000.0 for r in all_recipes)
    median = times[len(times) // 2]
    if MIN_PLAUSIBLE_ACTION_SEC <= median <= MAX_PLAUSIBLE_ACTION_SEC:
        return
    print(f"⚠ Median-Aktionszeit liegt bei {median:.4f}s (erwartet {MIN_PLAUSIBLE_ACTION_SEC}-"
          f"{MAX_PLAUSIBLE_ACTION_SEC}s). Das Script setzt voraus, dass 'BaseTime' in "
          f"MILLISEKUNDEN geliefert wird - stimmt das nicht mehr, sind ALLE Gold/h- und "
          f"XP/h-Werte um denselben Faktor falsch. Einheit im API-Feld pruefen.")


def build_recipe_by_output(all_recipes: list) -> dict:
    """1 Eintrag pro ItemId: das Recipe mit der kuerzesten Zeit/Stueck (Standardweg fuer
    die Kettenaufloesung, falls mehrere Skills dasselbe Item produzieren). Zeit haengt
    nicht vom Best-/Worst-Case ab, daher liefert dies fuer beide Cases dieselbe Auswahl.

    BEWUSSTE VEREINFACHUNG: ausgewaehlt wird rein nach Zeit, nicht nach Gewinn. Gibt es
    fuer ein Item einen schnellen-aber-teuren und einen langsamen-aber-guenstigen Weg,
    rechnet die Ketten-Analyse mit dem schnellen - der kann in Summe schlechter sein.
    Betrifft nur Items mit mehreren Herstellwegen; im Rohdaten-Tab sind beide Wege
    weiterhin einzeln sichtbar."""
    best = {}
    for r in all_recipes:
        time_per_unit = r["base_time_ms"] / r["item_amount"]
        current = best.get(r["item_id"])
        if current is None or time_per_unit < current["base_time_ms"] / current["item_amount"]:
            best[r["item_id"]] = r
    return best


def build_fish_to_cooked_map(all_recipes: list, recipe_by_output: dict) -> dict:
    """RawFishId -> CookedFishId, fuer Recipes mit genau 1 Zutat aus Fishing (Auto-Cook-Upgrade).
    Unabhaengig vom Best-/Worst-Case (betrifft nur Anzahl/Struktur der Cost-Zeilen, nicht
    deren Menge) - kann mit den Recipes aus einem beliebigen Case gebaut werden."""
    fishing_ids = {iid for iid, r in recipe_by_output.items() if r["skill"] == AUTO_COOK_SOURCE_SKILL}
    fish_to_cooked = {}
    for r in all_recipes:
        if r["skill"] == "Cooking" and len(r["costs"]) == 1:
            raw_id = r["costs"][0]["Item"]
            if raw_id in fishing_ids:
                fish_to_cooked[raw_id] = r["item_id"]
    return fish_to_cooked


# ============================================================
# 7. EINZELSCHRITT-ANALYSE ("Rohdaten"-Tab)
# ============================================================

def build_single_step_df(all_recipes: list, market_map: dict, item_info_map: dict) -> pd.DataFrame:
    results = []
    for r in all_recipes:
        if not is_player_shop_tradeable(item_info_map.get(r["item_id"], {})):
            continue

        sell_price, sold_to_npc = effective_sell_price(r["item_id"], market_map, item_info_map)

        # m wird fuer Zusatzinfos (Avg-Preis, Volumen, Anomalie-Checks) weiterverwendet.
        # Bei sold_to_npc=True ist der Player-Markt entweder nicht liquide genug oder
        # schlechter als der NPC-Preis - Markt-Warnungen/Volumen sind dann irrelevant.
        m = market_map.get(r["item_id"]) or {"buy": 0, "sell": 0, "buyVol": 0, "sellVol": 0, "avg": 0}

        # v10: KEIN Hard-Filter mehr hier - jedes Recipe landet im Rohdaten-Tab, damit man
        # sieht, WESHALB ein Item in Ketten/Realistisch_Farmbar fehlt (dort gilt weiterhin
        # exakt dieselbe Bedingung: sell_price<=0 -> ausgeschlossen, s. build_chain_df).
        # Grund wird unten in "AusschlussGrund" als Klartext festgehalten und im Excel-
        # Export rot/orange markiert statt die Zeile stillschweigend zu verwerfen.
        #
        # ACHTUNG zur Lesart: "InKetten" heisst genau "das Item ist ueberhaupt verkaufbar".
        # Ein Item kann InKetten=True haben und trotzdem NICHT im Ketten-Tab auftauchen,
        # naemlich wenn ein anderes (schnelleres) Rezept desselben Items den Platz belegt -
        # build_chain_df kennt pro ItemId nur einen Weg (s. build_recipe_by_output).
        in_ketten = sell_price > 0
        exclusion_reason = ""
        if not in_ketten:
            if m["buy"] <= 0 or m["sell"] <= 0:
                exclusion_reason = "Kein Markteintrag"
            elif not valid_market(m):
                exclusion_reason = f"Markt zu duenn (Volumen < {MIN_MARKET_VOLUME})"
            elif m["buyVol"] < MIN_SELL_VOLUME:
                exclusion_reason = f"Zu wenig Nachfrage (BuyVol {m['buyVol']:,} < {MIN_SELL_VOLUME:,})"
            else:
                exclusion_reason = "Kein Verkaufspreis"
            exclusion_reason += " + kein NPC-Verkauf moeglich"  # bei sell_price<=0 immer der Fall

        actions_per_hour = 3_600_000.0 / r["base_time_ms"]
        items_per_hour = actions_per_hour * r["item_amount"]
        gold_per_hour = items_per_hour * sell_price
        xp_per_hour = actions_per_hour * r["xp"]

        # Ø-Preis-Variante: dailyAveragePrice statt Bid, fuer Sell-Order-Strategie statt
        # Sofortverkauf (realistischer bei breitem Spread, siehe SpreadWarning). Nur
        # relevant, wenn ueberhaupt am Player-Markt verkauft wird.
        revenue_per_hour_avg = items_per_hour * m["avg"] if (not sold_to_npc and m["avg"] > 0) else None

        cost_per_action = 0.0
        cost_data_complete = True
        ingredient_market_unhealthy = False
        ingredient_price_anomaly = False
        ingredient_spread_warning = False
        max_liquidity_ratio = 0.0

        for c in r["costs"]:
            cm = market_map.get(c["Item"])
            if cm is None:
                cost_data_complete = False
                ingredient_market_unhealthy = True
                continue
            cost_per_action += cm["sell"] * c["Amount"]
            if not valid_market(cm):
                ingredient_market_unhealthy = True
            if price_anomaly(cm):
                ingredient_price_anomaly = True
            if wide_spread(cm):
                ingredient_spread_warning = True
            if cm["sellVol"] > 0:
                max_liquidity_ratio = max(max_liquidity_ratio, (c["Amount"] * actions_per_hour) / cm["sellVol"])

        if not sold_to_npc and m["buyVol"] > 0:
            max_liquidity_ratio = max(max_liquidity_ratio, items_per_hour / m["buyVol"])

        cost_per_hour = cost_per_action * actions_per_hour
        profit_per_hour = gold_per_hour - cost_per_hour
        profit_per_hour_avg = (revenue_per_hour_avg - cost_per_hour) if revenue_per_hour_avg is not None else None
        output_anomaly = price_anomaly(m) if not sold_to_npc else False
        output_spread_warning = wide_spread(m) if not sold_to_npc else False

        liquidity_warning = max_liquidity_ratio > LIQUIDITY_WARNING_RATIO or ingredient_market_unhealthy
        price_anomaly_flag = output_anomaly or ingredient_price_anomaly
        spread_warning_flag = output_spread_warning or ingredient_spread_warning

        if in_ketten:
            status = "OK"
            warn_parts = []
            if liquidity_warning:
                warn_parts.append("Liquiditaet")
            if price_anomaly_flag:
                warn_parts.append("Preisanomalie")
            if spread_warning_flag:
                warn_parts.append("Spread")
            if warn_parts:
                status = "Warnung"
                exclusion_reason = "Warnung: " + ", ".join(warn_parts)
        else:
            status = "Ausgeschlossen"

        results.append({
            "Item": r["name"], "Skill": r["skill"], "Level": r["level"],
            "Time_sec": r["base_time_ms"] / 1000.0,
            "XP": r["xp"],
            "Gold/h": profit_per_hour,
            "Gold/h (Ø-Preis)": profit_per_hour_avg,
            "Gold pro Stück": (profit_per_hour / items_per_hour) if items_per_hour > 0 else 0.0,
            "Stück/h": items_per_hour,
            "Revenue/h": gold_per_hour,
            "Revenue/h (Ø-Preis)": revenue_per_hour_avg,
            "Cost/h": cost_per_hour,
            "SoldToNPC": sold_to_npc,
            "NPCPreis": npc_sell_price(r["item_id"], item_info_map),
            "MarketAsk": m["sell"],
            "CostDataComplete": cost_data_complete,
            "InKetten": in_ketten,
            "Status": status,
            "AusschlussGrund": exclusion_reason,
            "LiquidityWarning": liquidity_warning,
            "LiquidityRatio": round(max_liquidity_ratio, 1),
            "IngredientMarketUnhealthy": ingredient_market_unhealthy,
            "PriceAnomaly": price_anomaly_flag,
            "OutputPriceAnomaly": output_anomaly,
            "IngredientPriceAnomaly": ingredient_price_anomaly,
            "SpreadWarning": spread_warning_flag,
            "OutputSpreadWarning": output_spread_warning,
            "IngredientSpreadWarning": ingredient_spread_warning,
            "BuyVol": m["buyVol"], "SellVol": m["sellVol"],
            "XP/h": xp_per_hour,
            "Gold per XP": profit_per_hour / max(xp_per_hour, 1),
            "ItemID": r["item_id"], "TaskId": r["task_id"],
            "CostCaseAmbiguous": r["cost_case_ambiguous"],
        })
    return pd.DataFrame(results)


def _merge_worst_case(df_best: pd.DataFrame, df_worst: pd.DataFrame,
                       keys: list[str], worst_cols: list[str], label: str) -> pd.DataFrame:
    """Gemeinsamer Merge-Kern fuer Rohdaten und Ketten.

    how="left" + validate="one_to_one": Best-Case ist die fuehrende Tabelle, es duerfen
    keine Zeilen dazukommen oder verschwinden. Ein Inner-Join haette bei doppelten
    Schluesseln (z.B. zwei Rezepte mit derselben ItemID und TaskId=None - pandas
    behandelt NaN-Schluessel beim Merge als gleich) stillschweigend Zeilen dupliziert
    und damit jede Auswertung darueber verfaelscht."""
    dup_best = df_best.duplicated(subset=keys).sum()
    dup_worst = df_worst.duplicated(subset=keys).sum()
    if dup_best or dup_worst:
        print(f"⚠ {label}: {dup_best} doppelte Schluessel im Best-Case, {dup_worst} im Worst-Case "
              f"(Schluessel: {keys}). Worst-Case-Spalten werden uebersprungen, damit keine "
              f"Zeilen vervielfacht werden - bitte Rezeptdaten pruefen.")
        return df_best
    return df_best.merge(df_worst[keys + worst_cols], on=keys, how="left",
                         suffixes=("", "_Worst"), validate="one_to_one")


def merge_worst_case_single_step(df_best: pd.DataFrame, df_worst: pd.DataFrame) -> pd.DataFrame:
    """Fuegt dem Rohdaten-Df die Worst-Case-Spalten hinzu (Smelting-Magic-Reichweite).
    Zeit/Stueck/XP sind vom Cost-Case unabhaengig - nur Cost/h, Gold/h, Gold pro Stück
    und Gold per XP koennen abweichen (nur bei *_bar-Rezepten mit >1 Cost-Zeile;
    fuer alle anderen Zeilen ist Best==Worst)."""
    if df_best.empty or df_worst.empty:
        return df_best
    return _merge_worst_case(
        df_best, df_worst,
        keys=["ItemID", "TaskId"],
        worst_cols=["Cost/h", "Gold/h", "Gold pro Stück", "Gold per XP"],
        label="Rohdaten-Worst-Case",
    )


# ============================================================
# 8. KETTEN-ANALYSE ("Ketten" / "Realistisch_Farmbar"-Tabs)
# ============================================================

def resolve_chain(item_id, market_map, recipe_by_output, fish_to_cooked,
                   qty_needed=1.0, visited=None, depth=0, max_depth=15):
    """Loest rekursiv auf, wie man `qty_needed` Stueck von item_id selbst herstellt statt
    am Markt zu kaufen. Beruecksichtigt das Auto-Cook-Upgrade (Fishing -> Cooking). Die
    Best-/Worst-Case-Unterscheidung steckt bereits in recipe_by_output["costs"] (s.
    normalize_recipe) - diese Funktion selbst ist Case-agnostisch.

    Gibt zurueck: (total_time_ms, total_gold_cost, steps, max_liquidity_ratio, fully_self_sufficient)
    """
    if visited is None:
        visited = set()

    if item_id in visited or depth > max_depth:
        price = market_map.get(item_id, {}).get("sell", 0)
        return 0.0, price * qty_needed, [], 0.0, False

    # Auto-Cook-Fall: item_id ist ein gekochter Fisch, der zu AUTO_COOK_CHANCE gratis mitkommt.
    # BEWUSST KONSERVATIV: die restlichen (1 - AUTO_COOK_CHANCE) rohen Fische werden weder
    # verkauft noch nachtraeglich gekocht - der Wert dieses "Beifangs" faellt unter den
    # Tisch. Die Kette rechnet sich real also eher besser als hier ausgewiesen.
    fish_source_id = next((raw for raw, cooked in fish_to_cooked.items() if cooked == item_id), None)
    if fish_source_id is not None and fish_source_id not in visited:
        fish_recipe = recipe_by_output[fish_source_id]
        visited2 = visited | {item_id, fish_source_id}

        cooked_per_action = fish_recipe["item_amount"] * AUTO_COOK_CHANCE
        actions_needed = qty_needed / cooked_per_action
        fishing_time_ms = actions_needed * fish_recipe["base_time_ms"]

        total_time_ms, total_cost, max_ratio = fishing_time_ms, 0.0, 0.0
        fully_self_sufficient = True
        steps = [(fish_recipe["name"] + " (mit Auto-Cook)", fish_recipe["skill"], qty_needed, fishing_time_ms)]

        for c in fish_recipe["costs"]:
            qty = c["Amount"] * actions_needed
            sub_time, sub_cost, sub_steps, sub_ratio, sub_ok = resolve_chain(
                c["Item"], market_map, recipe_by_output, fish_to_cooked, qty, visited2, depth + 1, max_depth
            )
            total_time_ms += sub_time
            total_cost += sub_cost
            max_ratio = max(max_ratio, sub_ratio)
            fully_self_sufficient = fully_self_sufficient and sub_ok
            steps.extend(sub_steps)

        return total_time_ms, total_cost, steps, max_ratio, fully_self_sufficient

    # Normalfall: kein eigenes Recipe -> am Markt kaufen
    if item_id not in recipe_by_output:
        market_entry = market_map.get(item_id)
        price = market_entry.get("sell", 0) if market_entry else 0
        sell_vol = market_entry.get("sellVol", 0) if market_entry else 0
        ratio = (qty_needed / sell_vol) if sell_vol > 0 else 0.0
        if market_entry is None or not valid_market(market_entry):
            ratio = max(ratio, 999.0)  # erzwingt LiquidityWarning, Markt zu unzuverlaessig
        return 0.0, price * qty_needed, [], ratio, False

    recipe = recipe_by_output[item_id]
    visited = visited | {item_id}
    actions_needed = qty_needed / recipe["item_amount"]
    own_time_ms = actions_needed * recipe["base_time_ms"]

    total_time_ms, total_cost, max_ratio = own_time_ms, 0.0, 0.0
    fully_self_sufficient = True
    steps = [(recipe["name"], recipe["skill"], qty_needed, own_time_ms)]

    for c in recipe["costs"]:
        qty = c["Amount"] * actions_needed
        sub_time, sub_cost, sub_steps, sub_ratio, sub_ok = resolve_chain(
            c["Item"], market_map, recipe_by_output, fish_to_cooked, qty, visited, depth + 1, max_depth
        )
        total_time_ms += sub_time
        total_cost += sub_cost
        max_ratio = max(max_ratio, sub_ratio)
        fully_self_sufficient = fully_self_sufficient and sub_ok
        steps.extend(sub_steps)

    return total_time_ms, total_cost, steps, max_ratio, fully_self_sufficient


def build_chain_df(recipe_by_output: dict, market_map: dict, item_info_map: dict, fish_to_cooked: dict) -> pd.DataFrame:
    chain_results = []
    for item_id, recipe in recipe_by_output.items():
        if not is_player_shop_tradeable(item_info_map.get(item_id, {})):
            continue

        sell_price, sold_to_npc = effective_sell_price(item_id, market_map, item_info_map)
        if sell_price <= 0:
            continue

        m = market_map.get(item_id) or {"buy": 0, "sell": 0, "buyVol": 0, "sellVol": 0, "avg": 0}
        if not sold_to_npc and m["buyVol"] < MIN_SELL_VOLUME:
            continue

        total_time_ms, raw_cost, steps, raw_ratio, fully_self_sufficient = resolve_chain(
            item_id, market_map, recipe_by_output, fish_to_cooked
        )
        if total_time_ms <= 0:
            continue

        actions_per_hour = 3_600_000.0 / total_time_ms
        revenue_per_hour = actions_per_hour * sell_price
        cost_per_hour = actions_per_hour * raw_cost
        profit_per_hour = revenue_per_hour - cost_per_hour

        # Ø-Preis-Variante (siehe Kommentar in build_single_step_df)
        revenue_per_hour_avg = actions_per_hour * m["avg"] if (not sold_to_npc and m["avg"] > 0) else None
        profit_per_hour_avg = (revenue_per_hour_avg - cost_per_hour) if revenue_per_hour_avg is not None else None

        max_liquidity_ratio = raw_ratio * actions_per_hour
        if not sold_to_npc and m["buyVol"] > 0:
            max_liquidity_ratio = max(max_liquidity_ratio, actions_per_hour / m["buyVol"])

        skills_involved = sorted(set(s[1] for s in steps))
        xp_per_unit_final_step = recipe["xp"] / recipe["item_amount"]

        chain_results.append({
            "Item": recipe["name"], "ItemID": item_id, "Level": recipe["level"],
            "FinalSkill": recipe["skill"],
            "ChainSkills": " -> ".join(skills_involved),
            "ChainDepth": len(steps),
            "FullySelfSufficient": fully_self_sufficient,
            "TimePerItem_sec": total_time_ms / 1000.0,
            "XP_letzter_Schritt_pro_Stück": xp_per_unit_final_step,
            "Gold/h (Eigenherstellung)": profit_per_hour,
            "Gold/h (Eigenherstellung, Ø-Preis)": profit_per_hour_avg,
            "Gold pro Stück": profit_per_hour / actions_per_hour,
            "Stück/h": actions_per_hour,
            "Revenue/h": revenue_per_hour,
            "Revenue/h (Ø-Preis)": revenue_per_hour_avg,
            "RawMaterialCost/h": cost_per_hour,
            "SoldToNPC": sold_to_npc,
            "NPCPreis": npc_sell_price(item_id, item_info_map),
            "MarketAsk": m["sell"],
            "LiquidityWarning": max_liquidity_ratio > LIQUIDITY_WARNING_RATIO,
            "LiquidityRatio": round(max_liquidity_ratio, 1),
            "PriceAnomaly": price_anomaly(m) if not sold_to_npc else False,
            "SpreadWarning": wide_spread(m) if not sold_to_npc else False,
            "XP/h (letzter Schritt)": xp_per_unit_final_step * actions_per_hour,
        })
    return pd.DataFrame(chain_results)


def merge_worst_case_chain(df_chain_best: pd.DataFrame, df_chain_worst: pd.DataFrame) -> pd.DataFrame:
    """Analog zu merge_worst_case_single_step, aber fuer die Ketten-Analyse: hier wirkt
    sich die Smelting-Magic-Reichweite auch auf alle nachgelagerten Items aus (z.B.
    haengt titanium_platebody an titanium_bar -> auch dessen Kette bekommt Worst-Spalten,
    obwohl das Platebody-Rezept selbst kein *_bar-Rezept ist)."""
    if df_chain_best.empty or df_chain_worst.empty:
        return df_chain_best
    merged = _merge_worst_case(
        df_chain_best, df_chain_worst,
        keys=["ItemID"],
        worst_cols=[
            "TimePerItem_sec", "Stück/h",
            "Gold/h (Eigenherstellung)", "Gold/h (Eigenherstellung, Ø-Preis)",
            "Gold pro Stück", "XP/h (letzter Schritt)",
        ],
        label="Ketten-Worst-Case",
    )
    if "Gold/h (Eigenherstellung)_Worst" not in merged.columns:
        return merged  # Merge wurde wegen doppelter Schluessel uebersprungen
    merged["CostCaseAmbiguous"] = (
        merged["Gold/h (Eigenherstellung)"] - merged["Gold/h (Eigenherstellung)_Worst"]
    ).abs() > 1e-6
    return merged


def fetch_orderbook_depth(item_id: int) -> dict | None:
    try:
        r = requests.get(COMPREHENSIVE_URL_TEMPLATE.format(item_id=item_id), timeout=15)
        return r.json() if r.status_code == 200 else None
    except Exception:
        return None


def enrich_with_longterm_averages(df: pd.DataFrame) -> pd.DataFrame:
    """Ergaenzt Rohdaten um Avg1D/Avg7D/Avg30D/Volume1D aus dem comprehensive-Endpoint -
    1 Request pro EINDEUTIGEM ItemID (dedupliziert, da dasselbe Item ueber mehrere
    Rezepte/Skills auftauchen kann). Nur aktiv bei SHOW_LONGTERM_AVERAGES=True, da das
    bei vielen Items in Rohdaten mehrere Minuten dauern kann. Nur fuer Rohdaten gedacht -
    Ketten/Realistisch_Farmbar/Nach_Skill_Level bleiben bewusst schlank (s. Absprache)."""
    unique_ids = [int(x) for x in df["ItemID"].unique().tolist()]
    print(f"Hole 1-/7-/30-Tage-Durchschnitt fuer {len(unique_ids)} eindeutige Items "
          f"(comprehensive-Endpoint, 1 Request/Item, das dauert einen Moment)...")

    rows = []
    fields_missing_warned = False
    for i, item_id in enumerate(unique_ids, 1):
        depth = fetch_orderbook_depth(item_id)
        entry = {"ItemID": item_id}
        if depth is not None:
            for col, field in COMPREHENSIVE_AVG_FIELDS.items():
                entry[col] = depth.get(field)
            entry["Volume1D"] = depth.get(COMPREHENSIVE_VOLUME_FIELD)
            if not fields_missing_warned and all(entry[c] is None for c in COMPREHENSIVE_AVG_FIELDS):
                print(f"⚠ Erwartete Avg-Felder {list(COMPREHENSIVE_AVG_FIELDS.values())} nicht in der "
                      f"Antwort gefunden. Tatsaechliche Keys: {sorted(depth.keys())}\n"
                      f"  -> Feldnamen in COMPREHENSIVE_AVG_FIELDS anpassen.")
                fields_missing_warned = True
        else:
            for col in COMPREHENSIVE_AVG_FIELDS:
                entry[col] = None
            entry["Volume1D"] = None
        rows.append(entry)

        if i % 25 == 0 or i == len(unique_ids):
            print(f"  ... {i}/{len(unique_ids)}")
        time.sleep(LONGTERM_AVERAGES_REQUEST_DELAY_S)

    avg_df = pd.DataFrame(rows)
    return df.merge(avg_df, on="ItemID", how="left")


def print_orderbook_depth(df_chain: pd.DataFrame):
    """Optionale Live-Tiefenpruefung (5 Buy/Sell-Stufen) fuer die Top-Selbstversorger-
    Items, zum direkten Abgleich mit dem Spiel-UI. Macht DEPTH_ANALYSIS_TOP_N zusaetzliche
    Requests - nur bei Bedarf einschalten (SHOW_ORDERBOOK_DEPTH = True)."""
    top = df_chain[df_chain["FullySelfSufficient"]].sort_values(
        "Gold/h (Eigenherstellung)", ascending=False
    ).head(DEPTH_ANALYSIS_TOP_N)

    print(f"\n=== ORDERBOOK-TIEFE (Top-{DEPTH_ANALYSIS_TOP_N} selbst herstellbar) ===")
    for _, row in top.iterrows():
        data = fetch_orderbook_depth(int(row["ItemID"]))
        if data is None:
            continue
        buys = sorted([(e["key"], int(e["value"])) for e in data.get("highestBuyPricesWithVolume", [])],
                      key=lambda x: x[0], reverse=True)[:5]
        sells = sorted([(e["key"], int(e["value"])) for e in data.get("lowestSellPricesWithVolume", [])],
                       key=lambda x: x[0])[:5]
        print(f"\n{row['Item']} - Gold/h: {row['Gold/h (Eigenherstellung)']:,.0f} ({row['FinalSkill']})")
        for i in range(max(len(buys), len(sells))):
            b = f"{buys[i][1]:>10,} @ {buys[i][0]:>6,}g" if i < len(buys) else ""
            s = f"{sells[i][1]:>10,} @ {sells[i][0]:>6,}g" if i < len(sells) else ""
            print(f"  BUY {b:30s}  SELL {s}")


def price_points_from_depth(depth: dict) -> list:
    """5 hoechste Buy-Preise (aufsteigend, niedrigster zuerst) + 5 niedrigste Sell-Preise
    (aufsteigend). Fehlende Tiefenstufen (dünner Markt) werden mit None aufgefuellt."""
    buys = sorted(depth.get("highestBuyPricesWithVolume", []), key=lambda x: x["key"], reverse=True)[:5]
    buys = sorted(buys, key=lambda x: x["key"])
    while len(buys) < 5:
        buys.insert(0, None)
    sells = sorted(depth.get("lowestSellPricesWithVolume", []), key=lambda x: x["key"])[:5]
    while len(sells) < 5:
        sells.append(None)
    return [(b["key"] if b else None) for b in buys] + [(s["key"] if s else None) for s in sells]


PRICE_SENSITIVITY_LABELS_SHORT = ["Buy1", "Buy2", "Buy3", "Buy4", "Buy5", "Sell1", "Sell2", "Sell3", "Sell4", "Sell5"]


def build_price_sensitivity_data(df_chain: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Holt einmalig die Live-Orderbook-Tiefe fuer die Top-N FullySelfSufficient-Items
    (macht PRICE_SENSITIVITY_TOP_N Requests) und gibt sie tabellarisch zurueck: 1 Zeile
    pro Item, je Preispunkt (Buy1..Sell5) zwei Spalten Preis + Gold/s. Wird sowohl fuer
    den PNG-Chart als auch fuer das Excel-Sheet "Preis_Sensitivitaet" verwendet, damit
    nicht zweimal live abgefragt wird. Gibt zusaetzlich die Liste der ueber NPC statt
    Player-Markt verkauften Items zurueck (fixer Preis, nicht Teil der Orderbook-Kurve).

    Die *_Gold_s-Spalten sind NETTO (Materialkosten der Kette schon abgezogen), damit sie
    direkt mit "Gold/h (Eigenherstellung)" / 3600 vergleichbar sind. Bis v10 stand hier
    der Brutto-Umsatz - dieselbe Achse, aber eine andere Groesse als ueberall sonst."""
    top = df_chain[df_chain["FullySelfSufficient"]].sort_values(
        "Gold/h (Eigenherstellung)", ascending=False
    ).head(PRICE_SENSITIVITY_TOP_N)

    rows = []
    npc_items = []
    for _, row in top.iterrows():
        depth = fetch_orderbook_depth(int(row["ItemID"]))
        if depth is None:
            continue
        prices = price_points_from_depth(depth)
        items_per_sec = row["Stück/h"] / 3600.0
        cost_per_item = (row["RawMaterialCost/h"] / row["Stück/h"]) if row["Stück/h"] > 0 else 0.0

        data = {
            "Item": row["Item"], "ItemID": int(row["ItemID"]), "FinalSkill": row["FinalSkill"],
            "SoldToNPC": bool(row["SoldToNPC"]),
            "Kosten_pro_Stück": cost_per_item,
        }
        for label, p in zip(PRICE_SENSITIVITY_LABELS_SHORT, prices):
            data[f"{label}_Preis"] = p
            data[f"{label}_Gold_s"] = (items_per_sec * (p - cost_per_item)) if p is not None else None
        rows.append(data)

        if bool(row["SoldToNPC"]):
            npc_items.append(str(row["Item"]))

    return pd.DataFrame(rows), npc_items


def build_price_sensitivity_chart(df_sens: pd.DataFrame, npc_items: list[str],
                                   path: str = PRICE_SENSITIVITY_CHART_PATH):
    """Matplotlib-Liniendiagramm: Gold/s NETTO (Best-Case-Kette) ueber die 10 aktuellen
    Orderbook-Preispunkte, eine Linie pro Top-N-Item, gemeinsame Y-Achse. Nutzt die
    bereits von build_price_sensitivity_data() geholten Daten (keine eigenen Requests).

    matplotlib ist eine reine Komfort-Abhaengigkeit: fehlt sie, gibt es hier einen Hinweis
    statt eines Absturzes - der Excel-Export (inkl. Sheet "Preis_Sensitivitaet" mit
    denselben Zahlen) laeuft dann ganz normal weiter."""
    try:
        import matplotlib.pyplot as plt  # lazy: nur noetig, wenn Flag aktiv
    except ImportError:
        print("ℹ matplotlib nicht installiert - Preis-Sensitivitaets-Chart wird uebersprungen "
              "(die Zahlen stehen trotzdem im Excel-Sheet 'Preis_Sensitivitaet'). "
              "Nachinstallieren mit: pip install matplotlib")
        return

    if df_sens.empty:
        print("Preis-Sensitivitaets-Chart: keine FullySelfSufficient-Items gefunden.")
        return

    labels = ["Buy1\n(niedrigster)", "Buy2", "Buy3", "Buy4", "Buy5\n(höchster Bid)",
              "Sell1\n(billigster)", "Sell2", "Sell3", "Sell4", "Sell5\n(teuerster)"]

    fig, ax = plt.subplots(figsize=(12, 7))
    for _, row in df_sens.iterrows():
        ys_raw = [row[f"{label}_Gold_s"] for label in PRICE_SENSITIVITY_LABELS_SHORT]
        xs = [i for i, v in enumerate(ys_raw) if pd.notna(v)]
        ys = [v for v in ys_raw if pd.notna(v)]
        ax.plot(xs, ys, marker="o", label=str(row["Item"]))

    ax.set_xticks(range(10))
    ax.set_xticklabels(labels)
    ax.axvline(4.5, color="gray", linestyle="--", linewidth=1)
    ax.axhline(0, color="black", linewidth=0.8)
    ax.set_ylabel("Gold/s (netto, Materialkosten abgezogen)")
    ax.set_title(f"Top-{PRICE_SENSITIVITY_TOP_N} FullySelfSufficient: Gold/s über Orderbook-Tiefe (Best-Case-Kette)")
    ax.legend(loc="best", fontsize=8, ncol=2)
    ax.grid(True, alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)

    print(f"Preis-Sensitivitaets-Chart gespeichert: {path}")
    if npc_items:
        print(f"ℹ Aktuell ueber NPC-Vendor verkauft (fixer Preis, nicht auf der Kurve): {', '.join(npc_items)}")


# ============================================================
# 9. EXCEL-EXPORT
# ============================================================

SORT_COLUMN_PER_SHEET = {
    "Rohdaten": "Gold/h",
    "Nach_Skill_Level": "Level",
    "Ketten": "Gold/h (Eigenherstellung)",
    "Realistisch_Farmbar": "Gold/h (Eigenherstellung)",
}

HIGHLIGHT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
HIGHLIGHT_HEADER_FILL = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
HIGHLIGHT_HEADER_FONT = Font(bold=True)

# Rohdaten (v10): kein Hard-Filter mehr - Items, die aus Ketten/Realistisch_Farmbar/
# Nach_Skill_Level rausfallen (die bleiben weiterhin gefiltert) bzw. eine Warnung haben,
# bleiben in Rohdaten sichtbar und werden markiert: die ganze Zeile kraeftig orange,
# und zusaetzlich das Feld, das den Ausschlag gab ("AusschlussGrund"), kraeftig rot.
ROW_FLAG_FILL = PatternFill(start_color="F6B26B", end_color="F6B26B", fill_type="solid")     # Zeile: kraeftiges Orange
REASON_FIELD_FILL = PatternFill(start_color="E06666", end_color="E06666", fill_type="solid")  # verantwortliches Feld: kraeftiges Rot

SKILL_LEVEL_COLUMNS = ["Skill", "Level", "Item", "Gold/h", "Gold/h_Worst", "Gold pro Stück", "Stück/h",
                        "XP/h", "SoldToNPC", "NPCPreis", "ItemID"]


def export_excel(df: pd.DataFrame, df_chain: pd.DataFrame, path: str, df_sensitivity: pd.DataFrame | None = None):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # Reihenfolge bewusst: Ketten -> Realistisch_Farmbar -> Nach_Skill_Level -> Rohdaten
        # -> Preis_Sensitivitaet (die tatsaechlich farmbaren/relevanten Sichten zuerst,
        # Rohdaten + Chart-Daten als Referenz zuletzt).
        if not df_chain.empty:
            df_chain.sort_values("Gold/h (Eigenherstellung)", ascending=False).to_excel(
                writer, sheet_name="Ketten", index=False
            )
            realistic = df_chain[df_chain["FullySelfSufficient"]].sort_values(
                "Gold/h (Eigenherstellung)", ascending=False
            )
            if not realistic.empty:
                realistic.to_excel(writer, sheet_name="Realistisch_Farmbar", index=False)
        # Nach_Skill_Level bildet wie bisher nur tatsaechlich verkaufbare Items ab (InKetten=True).
        # Rohdaten (unten) bleibt die einzige vollstaendige, ungefilterte Sicht.
        skill_level_cols = [c for c in SKILL_LEVEL_COLUMNS if c in df.columns]
        df[df["InKetten"]].sort_values(["Skill", "Level"], na_position="last")[skill_level_cols].to_excel(
            writer, sheet_name="Nach_Skill_Level", index=False
        )
        df.sort_values("Gold/h", ascending=False).to_excel(writer, sheet_name="Rohdaten", index=False)
        if df_sensitivity is not None and not df_sensitivity.empty:
            df_sensitivity.to_excel(writer, sheet_name="Preis_Sensitivitaet", index=False)

        for sheet_name, sort_col in SORT_COLUMN_PER_SHEET.items():
            if sheet_name not in writer.sheets:
                continue
            ws = writer.sheets[sheet_name]
            header_cells = next(ws.iter_rows(min_row=1, max_row=1))
            col_idx = next((c.column for c in header_cells if c.value == sort_col), None)
            if col_idx is None:
                continue
            col_letter = get_column_letter(col_idx)
            ws[f"{col_letter}1"].fill = HIGHLIGHT_HEADER_FILL
            ws[f"{col_letter}1"].font = HIGHLIGHT_HEADER_FONT
            for row in range(2, ws.max_row + 1):
                ws[f"{col_letter}{row}"].fill = HIGHLIGHT_FILL

        # Rohdaten: jede Zeile mit Status != OK komplett orange, das verantwortliche Feld
        # ("AusschlussGrund") zusaetzlich rot. Gold/h-Zelle bleibt vom gelben Sortier-
        # Highlight unberuehrt (s.o.).
        if "Rohdaten" in writer.sheets:
            ws = writer.sheets["Rohdaten"]
            header_cells = next(ws.iter_rows(min_row=1, max_row=1))
            headers = {c.value: c.column for c in header_cells}
            status_col = headers.get("Status")
            gold_col = headers.get("Gold/h")
            grund_col = headers.get("AusschlussGrund")
            if status_col is not None:
                for row in range(2, ws.max_row + 1):
                    status = ws.cell(row=row, column=status_col).value
                    if status == "OK" or status is None:
                        continue
                    for col in range(1, ws.max_column + 1):
                        if col == gold_col:
                            continue
                        ws.cell(row=row, column=col).fill = ROW_FLAG_FILL
                    if grund_col is not None:
                        ws.cell(row=row, column=grund_col).fill = REASON_FIELD_FILL


# ============================================================
# 10. LAUF-SANITY-CHECK
#
# Vergleicht ein paar Kennzahlen mit dem letzten Lauf (run_stats_history.json), um
# API-/Schema-Aenderungen von echten Marktbewegungen zu unterscheiden. Struktur-Werte
# (Tasks/Rezepte/Items aus der Game-Data) sollten bei echten Content-Updates nur
# WACHSEN, nie schrumpfen - jeder Ruecko dort ist verdaechtig auf einen kaputten
# Parser/Endpoint. Markt-Werte duerfen organisch schwanken, ein ploetzlicher Einbruch
# ist trotzdem eine Pruefung wert (koennte auch eine kaputte Markt-Antwort sein statt
# eine echte Bewegung).
# ============================================================

RUN_STATS_PATH = "run_stats_history.json"
RUN_STATS_HISTORY_LIMIT = 50     # so viele Laeufe werden aufgehoben (Datei heisst schliesslich *_history)
MARKET_DROP_WARNING_RATIO = 0.3  # 30% Einbruch bei Markt-Kennzahlen wird gemeldet

STRUCTURAL_STATS = ["tasks_count", "recipes_count", "item_info_count"]
MARKET_STATS = ["market_map_count", "rohdaten_count", "in_ketten_count", "realistic_count"]


def compute_run_stats(tasks: dict, all_recipes: list, item_info_map: dict,
                       market_map: dict, df: pd.DataFrame, df_chain: pd.DataFrame) -> dict:
    return {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "tasks_count": len(tasks),
        "recipes_count": len(all_recipes),
        "item_info_count": len(item_info_map),
        "market_map_count": len(market_map),
        "rohdaten_count": len(df),
        "in_ketten_count": int(df["InKetten"].sum()) if "InKetten" in df.columns else None,
        "realistic_count": (
            int(df_chain["FullySelfSufficient"].sum())
            if not df_chain.empty and "FullySelfSufficient" in df_chain.columns else 0
        ),
    }


def load_run_stats_history(path: str = RUN_STATS_PATH) -> list:
    """Liste vergangener Laeufe, aeltester zuerst. Alte Dateien enthalten nur ein einzelnes
    Dict (nur der letzte Lauf) - das wird transparent als Ein-Element-Historie gelesen."""
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return []
    if isinstance(data, dict):
        return [data]
    if isinstance(data, list):
        return [entry for entry in data if isinstance(entry, dict)]
    return []


def load_last_run_stats(path: str = RUN_STATS_PATH) -> dict | None:
    history = load_run_stats_history(path)
    return history[-1] if history else None


def save_run_stats(stats: dict, path: str = RUN_STATS_PATH):
    history = load_run_stats_history(path)
    history.append(stats)
    history = history[-RUN_STATS_HISTORY_LIMIT:]
    with open(path, "w", encoding="utf-8") as f:
        json.dump(history, f, indent=2, ensure_ascii=False)


def sanity_check_run_stats(current: dict, previous: dict | None) -> list[str]:
    """Vergleicht aktuelle Kennzahlen mit dem letzten Lauf. Gibt eine Liste von
    Warntexten zurueck (leer = alles unauffaellig)."""
    if previous is None:
        return []

    warnings = []
    for key in STRUCTURAL_STATS:
        old, new = previous.get(key), current.get(key)
        if old is None or new is None:
            continue
        if new < old:
            warnings.append(
                f"'{key}': {old} -> {new} (GESUNKEN). Struktur-Werte aus der Game-Data sollten "
                f"bei echten Updates nur wachsen - moeglicher Hinweis auf ein kaputtes API-Feld/Parsing."
            )

    for key in MARKET_STATS:
        old, new = previous.get(key), current.get(key)
        if old is None or new is None or old == 0:
            continue
        drop_ratio = (old - new) / old
        if drop_ratio > MARKET_DROP_WARNING_RATIO:
            warnings.append(
                f"'{key}': {old} -> {new} ({drop_ratio:.0%} Einbruch). Kann eine echte Marktbewegung "
                f"sein, oder eine kaputte/unvollstaendige API-Antwort - pruefen lohnt sich."
            )

    return warnings


def print_run_sanity_check(current: dict, previous: dict | None):
    if previous is None:
        print("Lauf-Sanity-Check: erster Lauf, noch keine Vergleichsbasis (wird jetzt gespeichert).")
        return

    warnings = sanity_check_run_stats(current, previous)
    if not warnings:
        print(f"Lauf-Sanity-Check: unauffaellig (letzter Lauf: {previous.get('timestamp', '?')}).")
        return

    print(f"\n⚠⚠ LAUF-SANITY-CHECK: {len(warnings)} Auffaelligkeit(en) ggue. letztem Lauf "
          f"({previous.get('timestamp', '?')}):")
    for w in warnings:
        print(f"  - {w}")
    print("  Falls das nicht durch ein bekanntes Game-Update erklaerbar ist: API-Antwort direkt "
          "im Browser gegenpruefen statt den Zahlen blind zu vertrauen.\n")


# ============================================================
# 11. MAIN
# ============================================================

def print_summary(df: pd.DataFrame, df_chain: pd.DataFrame):
    excluded_count = int((~df["InKetten"]).sum())
    print(f"Rohdaten: {len(df)} Einzelschritt-Recipes ({excluded_count} davon nicht in Ketten - Grund s. AusschlussGrund).")
    if not df_chain.empty:
        n_self_sufficient = int(df_chain["FullySelfSufficient"].sum())
        print(f"Ketten: {len(df_chain)} Endprodukte, davon {n_self_sufficient} 100% selbst herstellbar.")

    dupe_count = int(df.duplicated("Item", keep=False).sum())
    if dupe_count:
        print(f"ℹ {dupe_count} Zeilen mit doppeltem Item-Namen (unterschiedliche Skills/Recipes) - Details im Rohdaten-Tab.")

    npc_count = int(df["SoldToNPC"].sum())
    if npc_count:
        print(f"ℹ {npc_count} Items werden ueber NPC-Vendor statt Player-Markt verkauft (kein liquider Markt oder NPC-Preis besser).")

    ambiguous_count = int(df["CostCaseAmbiguous"].sum())
    if ambiguous_count:
        print(f"ℹ {ambiguous_count} Rezepte mit Best-/Worst-Case-Unterschied (Smelting-Magic-Reichweite) - s. _Worst-Spalten.")
    if not df_chain.empty and "CostCaseAmbiguous" in df_chain.columns:
        chain_ambiguous_count = int(df_chain["CostCaseAmbiguous"].sum())
        if chain_ambiguous_count:
            print(f"ℹ {chain_ambiguous_count} Ketten-Items mit Best-/Worst-Case-Unterschied (direkt/indirekt via *_bar-Zutat).")

    warn = df[df["LiquidityWarning"]]
    if not warn.empty:
        print(f"⚠ {len(warn)} Items mit LiquidityWarning (Bedarf/Absatz > {LIQUIDITY_WARNING_RATIO:.0f}x Marktvolumen).")

    anomalies = df[df["PriceAnomaly"]]
    if not anomalies.empty:
        print(f"⚠ {len(anomalies)} Items mit PriceAnomaly (>{int(MAX_AVG_DEVIATION_RATIO * 100)}% Abweichung vom 24h-Avg).")

    spread_warnings = df[df["SpreadWarning"]]
    if not spread_warnings.empty:
        print(f"ℹ {len(spread_warnings)} Items mit SpreadWarning (Ask >{int(MAX_SPREAD_RATIO * 100)}% ueber Bid - tiefer Markt, ggf. per Sell-Order statt Sofortverkauf handeln).")


def main():
    previous_run_stats = load_last_run_stats()

    try:
        market_map = load_market_map()
        game = load_game_data()
    except RuntimeError as exc:
        print(f"⚠ Abbruch: {exc}")
        return

    item_info_map = build_item_info_map(game)
    tasks = game.get("Tasks", {})
    if not tasks:
        print("⚠ Abbruch: Game-Data enthaelt keinen 'Tasks'-Block - Endpunkt/Schema pruefen.")
        return

    # Best-/Worst-Case parallel aufbauen (Smelting-Magic-Reichweite, s. Abschnitt 2/6)
    all_recipes_best = build_all_recipes(tasks, case="best")
    all_recipes_worst = build_all_recipes(tasks, case="worst")
    check_action_time_plausibility(all_recipes_best)

    recipe_by_output_best = build_recipe_by_output(all_recipes_best)
    recipe_by_output_worst = build_recipe_by_output(all_recipes_worst)
    fish_to_cooked = build_fish_to_cooked_map(all_recipes_best, recipe_by_output_best)

    df_best = build_single_step_df(all_recipes_best, market_map, item_info_map)
    if df_best.empty:
        print("Keine brauchbaren Farming-Daten gefunden.")
        return
    df_worst = build_single_step_df(all_recipes_worst, market_map, item_info_map)
    df = merge_worst_case_single_step(df_best, df_worst)
    if SHOW_LONGTERM_AVERAGES:
        df = enrich_with_longterm_averages(df)

    df_chain_best = build_chain_df(recipe_by_output_best, market_map, item_info_map, fish_to_cooked)
    df_chain_worst = build_chain_df(recipe_by_output_worst, market_map, item_info_map, fish_to_cooked)
    df_chain = merge_worst_case_chain(df_chain_best, df_chain_worst)

    current_run_stats = compute_run_stats(tasks, all_recipes_best, item_info_map, market_map, df, df_chain)
    print_run_sanity_check(current_run_stats, previous_run_stats)
    save_run_stats(current_run_stats)

    print_summary(df, df_chain)
    if SHOW_ORDERBOOK_DEPTH and not df_chain.empty:
        print_orderbook_depth(df_chain)

    df_sensitivity = pd.DataFrame()
    if SHOW_PRICE_SENSITIVITY_CHART and not df_chain.empty:
        df_sensitivity, npc_items = build_price_sensitivity_data(df_chain)
        build_price_sensitivity_chart(df_sensitivity, npc_items)

    export_path = EXPORT_PATH
    try:
        export_excel(df, df_chain, export_path, df_sensitivity)
    except PermissionError:
        print(f"⚠ '{export_path}' ist gesperrt (meist: Datei ist noch in Excel geoeffnet).")
        fallback_path = f"{EXPORT_PATH.removesuffix('.xlsx')}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        print(f"  Schliesse die Datei in Excel fuer den Standardnamen - speichere stattdessen unter: {fallback_path}")
        export_excel(df, df_chain, fallback_path, df_sensitivity)
        export_path = fallback_path
    sheets = "Ketten, Realistisch_Farmbar, Nach_Skill_Level, Rohdaten"
    if not df_sensitivity.empty:
        sheets += ", Preis_Sensitivitaet"
    print(f"\nExcel gespeichert: {export_path} (Sheets: {sheets})")


if __name__ == "__main__":
    main()
